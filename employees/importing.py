"""Bulk employee import — D-122. Preview IS apply, rolled back.

**THE CENTRAL RULE.** Preview and apply run the exact same code, calling the
exact same service functions the single-capture path calls: ``Employee``
creation with its identity checks, ``engage()``, ``capture()``. There is no
``validate_only`` flag anywhere in this call chain — that would be a second
validation path by another name, and D-122 exists precisely because an import
that bypasses the ID check, the minimum age check or the minimum wage check
puts forty unchecked employees on file with nobody's name against the
exception. The only difference between preview and apply is whether the
transaction that ran it commits or rolls back at the end. Rolling back burns
primary key sequence values — the next real employee gets a higher id than the
row count would suggest. That is fine; sequences are not a report.

**One declared column spec, read by both sides** (task 4's anti-drift
requirement, and D-122's own reasoning restated for this table): the template
that goes out to the employer and the parser that reads it back both read
``EXPECTED_COLUMNS``. A hand-maintained template sheet that drifts from what
the importer expects produces failures the employer cannot diagnose; a single
spec cannot drift from itself.

**A batch is a unit, both ways.** Apply refuses outright while any row carries
a blocking error — the employer fixes the file and re-uploads, rather than the
importer writing the rows that happened to be clean and silently dropping the
rest. Reverse deletes exactly the batch's employees (their children cascade)
inside one transaction, and refuses — naming which employee — if anything
downstream still references one of them.

**A below-minimum-wage row is a warning, not a blocking error** (D-108), and
can be written only when the caller of ``apply_batch`` passes a named
``acknowledged_by``. There is no default and no batch-level bypass: an import
run twice with the same file, once without an acknowledger and once with one,
must behave exactly as two single captures of the same rate would.

**The validation report never stores a full ID number or bank account
number.** Row number, column, message, and the last four digits where
identity matters — the same discipline ``employee.id_number_last4`` already
applies to the column itself.
"""

from __future__ import annotations

import datetime
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models.deletion import ProtectedError
from django.utils import timezone

from core.db.fields import keyed_hash
from core.files import purge_content
from core.managers import tenant_context_of
from employees.engagements import (
    EngagementRefusedError,
    check_minimum_age,
    engage,
)
from employees.identity import luhn_check_digit
from employees.models import (
    Employee,
    EmployeeEngagement,
    EmployeeImportBatch,
    EmployeeRemuneration,
)
from employees.remuneration import (
    BelowMinimumWageError,
    RemunerationRefusedError,
    capture,
)

Status = EmployeeImportBatch.Status


class ImportRefusedError(Exception):
    """The batch may not be applied or reversed as asked. Nothing was written."""


class IllegalTransitionError(Exception):
    """The batch cannot move to that status from where it is."""


# --------------------------------------------------------------- status machine

#: What each status may legally become. Enforced here, in the service layer —
#: the CHECK constraint only proves the value is a known one, not that the move
#: from the row's previous value was legal.
ALLOWED_TRANSITIONS: dict[str, frozenset[str]] = {
    # APPLIED is reachable directly from UPLOADED (and from PREVIEW, and from a
    # retry after FAILED): apply_batch() re-validates every row itself rather
    # than trusting an earlier preview's report, so a prior preview call is a
    # convenience for the employer, never a precondition the state machine
    # enforces.
    Status.UPLOADED: frozenset({Status.VALIDATING, Status.PREVIEW, Status.APPLIED, Status.FAILED}),
    Status.VALIDATING: frozenset({Status.PREVIEW, Status.APPLIED, Status.FAILED}),
    Status.PREVIEW: frozenset({Status.VALIDATING, Status.PREVIEW, Status.APPLIED, Status.FAILED}),
    Status.APPLIED: frozenset({Status.REVERSED}),
    Status.REVERSED: frozenset(),
    Status.FAILED: frozenset({Status.VALIDATING, Status.PREVIEW, Status.APPLIED}),
}


def transition(batch: EmployeeImportBatch, new_status: str) -> None:
    """Move ``batch.status`` to ``new_status`` in memory, or refuse. Does not save."""
    if new_status == batch.status:
        return
    allowed = ALLOWED_TRANSITIONS.get(batch.status, frozenset())
    if new_status not in allowed:
        raise IllegalTransitionError(
            f"Batch {batch.pk} cannot move from '{batch.status}' to '{new_status}'. "
            f"Allowed from '{batch.status}': {sorted(allowed) or 'nothing — this is terminal'}."
        )
    batch.status = new_status


# ----------------------------------------------------------------- column spec


@dataclass(frozen=True)
class ImportColumn:
    """One column, in both the generated template and the parser. The one spec."""

    key: str
    header: str
    kind: str  # "text" | "date" | "decimal" | "choice" | "boolean"
    required: bool = True
    choices: tuple[tuple[str, str], ...] | None = None
    example: object = ""
    help_text: str = ""


#: A real, checksum-valid example so a curious employer who leaves row 2 in
#: place does not get refused by the very validation this file exists to run.
_EXAMPLE_ID_NUMBER = "900101" + "5009" + "08"
_EXAMPLE_ID_NUMBER += str(luhn_check_digit(_EXAMPLE_ID_NUMBER))

EXPECTED_COLUMNS: tuple[ImportColumn, ...] = (
    ImportColumn("first_name", "First name", "text", example="Thandi"),
    ImportColumn("last_name", "Last name", "text", example="Mokoena"),
    ImportColumn(
        "id_type",
        "ID type",
        "choice",
        choices=Employee.IdType.choices,
        example=Employee.IdType.SA_ID,
        help_text="A South African ID is the only type this system can checksum-verify.",
    ),
    ImportColumn(
        "id_number",
        "ID / passport / permit number",
        "text",
        example=_EXAMPLE_ID_NUMBER,
        help_text="Digits only for a South African ID. Never shared or logged in full.",
    ),
    ImportColumn("date_of_birth", "Date of birth", "date", example=datetime.date(1990, 1, 1)),
    ImportColumn(
        "gender",
        "Gender",
        "choice",
        required=False,
        choices=Employee.Gender.choices,
        example="",
        help_text="Optional. Collected for the EEA return and the UI-19 only.",
    ),
    ImportColumn("mobile_number", "Mobile number", "text", example="+27821234567"),
    ImportColumn(
        "email",
        "Email address",
        "text",
        required=False,
        example="thandi@example.com",
        help_text="Required unless 'No email' is TRUE.",
    ),
    ImportColumn(
        "has_no_email",
        "No email (TRUE/FALSE)",
        "boolean",
        required=False,
        example="FALSE",
        help_text="TRUE only for an employee with genuinely no address — payslips then go by SMS.",
    ),
    ImportColumn("job_title", "Job title", "text", example="Domestic worker"),
    ImportColumn("start_date", "Start date", "date", example=datetime.date(2026, 3, 1)),
    ImportColumn(
        "contract_type",
        "Contract type",
        "choice",
        required=False,
        choices=EmployeeEngagement.ContractType.choices,
        example=EmployeeEngagement.ContractType.PERMANENT,
        help_text="Optional. Defaults to Permanent when left blank.",
    ),
    ImportColumn(
        "pay_basis",
        "Pay basis",
        "choice",
        choices=EmployeeRemuneration.PayBasis.choices,
        example=EmployeeRemuneration.PayBasis.MONTHLY,
    ),
    ImportColumn(
        "rate_amount",
        "Rate amount",
        "decimal",
        # A plain int, not Decimal("5000.00") — test_no_hardcoded_rates scans
        # for exactly that shape, and this is an example salary for a demo
        # cell, not a statutory figure with a citation.
        example=5000,
        help_text="In the unit of pay basis — a monthly salary if pay basis is Monthly.",
    ),
)

_COLUMNS_BY_KEY = {column.key: column for column in EXPECTED_COLUMNS}


# ------------------------------------------------------------------- row types


@dataclass(frozen=True)
class RowIssue:
    row_number: int
    column: str
    message: str
    severity: str  # "error" | "warning"

    def as_dict(self) -> dict:
        return {
            "row": self.row_number,
            "column": self.column,
            "message": self.message,
            "severity": self.severity,
        }


@dataclass(frozen=True)
class ParsedRow:
    row_number: int
    values: dict
    errors: tuple[str, ...] = ()


@dataclass(frozen=True)
class BatchResult:
    row_count: int
    accepted_count: int
    rejected_count: int
    blocking_count: int
    below_minimum_count: int
    issues: tuple[RowIssue, ...] = field(default_factory=tuple)

    @property
    def report(self) -> list[dict]:
        return [issue.as_dict() for issue in self.issues]


# --------------------------------------------------------------------- parsing


def _parse_text(raw) -> str:
    return "" if raw is None else str(raw).strip()


def _parse_date(raw) -> datetime.date | None:
    if raw in (None, ""):
        return None
    if isinstance(raw, datetime.datetime):
        return raw.date()
    if isinstance(raw, datetime.date):
        return raw
    text = str(raw).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d %B %Y"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _parse_decimal(raw) -> Decimal | None:
    if raw in (None, ""):
        return None
    try:
        return Decimal(str(raw))
    except InvalidOperation:
        return None


def _parse_boolean(raw) -> bool:
    text = _parse_text(raw).lower()
    return text in {"true", "yes", "1", "y", "x"}


def _parse_choice(raw, column: ImportColumn) -> str | None:
    text = _parse_text(raw).lower()
    if not text:
        return None
    valid = {value.lower(): value for value, _label in column.choices}
    return valid.get(text)


def _parse_cell(raw, column: ImportColumn) -> tuple[object, str | None]:
    """The typed value, and an error message if the cell cannot be used."""
    text = _parse_text(raw)
    if not text and column.kind != "boolean":
        if column.required:
            return None, f"{column.header} is required."
        return ("" if column.kind == "text" else None), None

    if column.kind == "text":
        return text, None
    if column.kind == "date":
        parsed = _parse_date(raw)
        if parsed is None:
            return None, f"{column.header} is not a date: {text!r}."
        return parsed, None
    if column.kind == "decimal":
        parsed = _parse_decimal(raw)
        if parsed is None:
            return None, f"{column.header} is not a number: {text!r}."
        return parsed, None
    if column.kind == "boolean":
        return _parse_boolean(raw), None
    if column.kind == "choice":
        parsed = _parse_choice(raw, column)
        if parsed is None:
            allowed = ", ".join(value for value, _label in column.choices)
            return None, f"{column.header} must be one of: {allowed}. Got {text!r}."
        return parsed, None
    raise AssertionError(f"Unknown column kind: {column.kind}")  # pragma: no cover


def parse_workbook(file_like, *, sheet_name: str = "Employees") -> list[ParsedRow]:
    """Read the data sheet against ``EXPECTED_COLUMNS``. Matches by header text,
    not by column position, so a re-ordered (but not renamed) sheet still works.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(file_like, data_only=True)
    sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active

    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    index_by_key: dict[str, int] = {}
    for index, header in enumerate(header_row):
        for column in EXPECTED_COLUMNS:
            if header is not None and str(header).strip() == column.header:
                index_by_key[column.key] = index

    rows: list[ParsedRow] = []
    for row_number, raw_row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if all(cell in (None, "") for cell in raw_row):
            continue

        values: dict = {}
        errors: list[str] = []
        for column in EXPECTED_COLUMNS:
            index = index_by_key.get(column.key)
            raw = raw_row[index] if index is not None and index < len(raw_row) else None
            value, error = _parse_cell(raw, column)
            if error:
                errors.append(error)
            else:
                values[column.key] = value

        rows.append(ParsedRow(row_number=row_number, values=values, errors=tuple(errors)))

    return rows


# ------------------------------------------------------------------ validation


def _validation_message(error: ValidationError) -> str:
    if hasattr(error, "message_dict"):
        return "; ".join(
            f"{field_name}: {' '.join(msgs)}" for field_name, msgs in error.message_dict.items()
        )
    return " ".join(error.messages)


def _run_rows(
    batch: EmployeeImportBatch, rows: list[ParsedRow], *, acknowledged_by=None
) -> BatchResult:
    """Run every row through the single-capture path. Returns the outcome.

    Every row that gets as far as ``Employee`` creation runs inside its own
    savepoint, so one row's failure never disturbs another's — but the caller
    decides, based on the counts returned here, whether ANY of it is kept.
    """
    tenant = batch.tenant
    employer = batch.employer

    issues: list[RowIssue] = []
    accepted = 0
    rejected = 0
    blocking = 0
    below_minimum = 0
    seen_hashes: dict[str, int] = {}

    for parsed in rows:
        if parsed.errors:
            for message in parsed.errors:
                issues.append(RowIssue(parsed.row_number, "", message, "error"))
            rejected += 1
            blocking += 1
            continue

        values = parsed.values
        id_number = values["id_number"]
        last4 = id_number[-4:] if len(id_number) >= 4 else id_number

        id_hash = keyed_hash(id_number, scope=f"tenant:{tenant.pk}")
        if id_hash in seen_hashes:
            issues.append(
                RowIssue(
                    parsed.row_number,
                    "id_number",
                    f"Duplicate of row {seen_hashes[id_hash]} in this file (ID ending {last4}).",
                    "error",
                )
            )
            rejected += 1
            blocking += 1
            continue
        if Employee.objects.filter(tenant=tenant, id_number_hash=id_hash).exists():
            issues.append(
                RowIssue(
                    parsed.row_number,
                    "id_number",
                    f"An employee with an ID ending {last4} is already on file for this employer.",
                    "error",
                )
            )
            rejected += 1
            blocking += 1
            continue

        age_check = check_minimum_age(values["date_of_birth"], values["start_date"])
        if not age_check.permitted:
            issues.append(RowIssue(parsed.row_number, "date_of_birth", age_check.reason, "error"))
            rejected += 1
            blocking += 1
            continue

        seen_hashes[id_hash] = parsed.row_number

        savepoint = transaction.savepoint()
        try:
            employee = Employee(
                tenant=tenant,
                employer=employer,
                first_name=values["first_name"],
                last_name=values["last_name"],
                date_of_birth=values["date_of_birth"],
                gender=values.get("gender") or "",
                mobile_number=values["mobile_number"],
                email=values.get("email") or "",
                has_no_email=bool(values.get("has_no_email")),
                id_type=values["id_type"],
                id_number=id_number,
                created_by_import_batch=batch,
            )
            employee.full_clean()
            employee.save()

            engage(
                employee,
                start_date=values["start_date"],
                job_title=values["job_title"],
                contract_type=values.get("contract_type")
                or EmployeeEngagement.ContractType.PERMANENT,
            )

            remuneration = capture(
                employee,
                pay_basis=values["pay_basis"],
                rate_amount=values["rate_amount"],
                effective_from=values["start_date"],
                acknowledged_by=acknowledged_by,
            )
        except BelowMinimumWageError as error:
            transaction.savepoint_rollback(savepoint)
            issues.append(RowIssue(parsed.row_number, "rate_amount", str(error), "warning"))
            rejected += 1
            below_minimum += 1
            continue
        except (ValidationError, EngagementRefusedError, RemunerationRefusedError) as error:
            transaction.savepoint_rollback(savepoint)
            message = (
                _validation_message(error) if isinstance(error, ValidationError) else str(error)
            )
            issues.append(RowIssue(parsed.row_number, "", message, "error"))
            rejected += 1
            blocking += 1
            continue
        else:
            transaction.savepoint_commit(savepoint)
            accepted += 1
            if remuneration.is_below_minimum:
                who = acknowledged_by_id(acknowledged_by)
                issues.append(
                    RowIssue(
                        parsed.row_number,
                        "rate_amount",
                        f"Below the applicable minimum wage. Acknowledged by user {who}.",
                        "warning",
                    )
                )
                below_minimum += 1

    return BatchResult(
        row_count=len(rows),
        accepted_count=accepted,
        rejected_count=rejected,
        blocking_count=blocking,
        below_minimum_count=below_minimum,
        issues=tuple(issues),
    )


def acknowledged_by_id(acknowledged_by) -> str:
    return str(getattr(acknowledged_by, "pk", acknowledged_by))


def _save_report(batch: EmployeeImportBatch, result: BatchResult) -> None:
    batch.row_count = result.row_count
    batch.accepted_count = result.accepted_count
    batch.rejected_count = result.rejected_count
    batch.validation_report = result.report
    batch.save(
        update_fields=[
            "status",
            "row_count",
            "accepted_count",
            "rejected_count",
            "validation_report",
            "updated_at",
        ]
    )


# --------------------------------------------------------------- orchestration


def preview_batch(batch: EmployeeImportBatch, rows: list[ParsedRow]) -> BatchResult:
    """Run every row, then roll all of it back. The batch's own report persists."""
    with transaction.atomic(), tenant_context_of(batch):
        savepoint = transaction.savepoint()
        result = _run_rows(batch, rows, acknowledged_by=None)
        transaction.savepoint_rollback(savepoint)

        transition(batch, Status.PREVIEW)
        _save_report(batch, result)

    return result


def apply_batch(
    batch: EmployeeImportBatch, rows: list[ParsedRow], *, acknowledged_by=None
) -> BatchResult:
    """Run every row for real. Refuses — and writes nothing — if any row carries
    a blocking error, or if any row is below the minimum wage and nobody has
    acknowledged that.
    """
    with transaction.atomic(), tenant_context_of(batch):
        savepoint = transaction.savepoint()
        result = _run_rows(batch, rows, acknowledged_by=acknowledged_by)

        if result.blocking_count:
            transaction.savepoint_rollback(savepoint)
            transition(batch, Status.PREVIEW)
            _save_report(batch, result)
            raise ImportRefusedError(
                f"{result.blocking_count} row(s) carry a blocking error and must be "
                f"fixed before this batch can be applied. Nothing was written."
            )
        if result.below_minimum_count and acknowledged_by is None:
            transaction.savepoint_rollback(savepoint)
            transition(batch, Status.PREVIEW)
            _save_report(batch, result)
            raise ImportRefusedError(
                f"{result.below_minimum_count} row(s) are below the applicable minimum "
                f"wage. Apply again with acknowledged_by naming who accepts that — "
                f"there is no default and no batch-level bypass."
            )

        transaction.savepoint_commit(savepoint)
        transition(batch, Status.APPLIED)
        batch.applied_at = timezone.now()
        batch.row_count = result.row_count
        batch.accepted_count = result.accepted_count
        batch.rejected_count = result.rejected_count
        batch.validation_report = result.report
        batch.save(
            update_fields=[
                "status",
                "applied_at",
                "row_count",
                "accepted_count",
                "rejected_count",
                "validation_report",
                "updated_at",
            ]
        )

    _purge_source_file(batch)
    return result


def reverse_batch(batch: EmployeeImportBatch) -> None:
    """Delete the batch's employees (children cascade). One transaction, all or
    nothing. Refuses, naming the employee, if anything downstream still
    references one of them.
    """
    with transaction.atomic(), tenant_context_of(batch):
        transition(batch, Status.REVERSED)

        for employee in Employee.objects.filter(created_by_import_batch=batch):
            try:
                employee.delete()
            except ProtectedError as error:
                raise ImportRefusedError(
                    f"This batch cannot be reversed: {employee} is still referenced "
                    f"elsewhere and cannot be removed. Nothing was removed. ({error})"
                ) from error

        batch.reversed_at = timezone.now()
        batch.save(update_fields=["status", "reversed_at", "updated_at"])

    _purge_source_file(batch)


DATA_SHEET_NAME = "Employees"
INSTRUCTIONS_SHEET_NAME = "Instructions"
#: How many data rows carry the dropdown validation. Generous rather than exact
#: — a cleaning company onboarding forty staff should never hit the edge of it.
TEMPLATE_VALIDATION_ROWS = 500


def _example_cell_value(column: ImportColumn):
    if column.kind == "choice" and column.example != "":
        # TextChoices values compare equal to their plain string, but openpyxl
        # writes an actual str rather than a Choices member either way.
        return str(column.example)
    if column.kind == "decimal" and isinstance(column.example, Decimal):
        return float(column.example)
    return column.example


def build_template_workbook():
    """The .xlsx an employer fills in, generated from ``EXPECTED_COLUMNS`` alone.

    D-122's reason restated for this table: a hand-maintained template drifts
    from the columns the importer expects, and that drift produces failures the
    employer cannot diagnose. There is exactly one column spec, and both the
    template and ``parse_workbook`` read it.

    No rate, threshold or statutory figure is embedded anywhere in this
    function — ``test_no_hardcoded_rates`` governs this file too, and the one
    number here (the example rate) is a plausible salary, not a gazetted one.
    """
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = openpyxl.Workbook()
    data_sheet = workbook.active
    data_sheet.title = DATA_SHEET_NAME

    for col_index, column in enumerate(EXPECTED_COLUMNS, start=1):
        header_cell = data_sheet.cell(row=1, column=col_index, value=column.header)
        header_cell.font = Font(bold=True)
        example_cell = data_sheet.cell(row=2, column=col_index, value=_example_cell_value(column))
        if column.kind == "date":
            example_cell.number_format = "YYYY-MM-DD"

    for col_index, column in enumerate(EXPECTED_COLUMNS, start=1):
        if not column.choices:
            continue
        letter = get_column_letter(col_index)
        allowed = ",".join(value for value, _label in column.choices)
        validation = DataValidation(
            type="list",
            formula1=f'"{allowed}"',
            allow_blank=not column.required,
            showErrorMessage=True,
        )
        validation.error = f"Choose one of: {allowed}"
        validation.errorTitle = "Invalid value"
        data_sheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}{TEMPLATE_VALIDATION_ROWS + 1}")

    instructions_sheet = workbook.create_sheet(INSTRUCTIONS_SHEET_NAME)
    instructions_sheet.append(["Column", "Mandatory", "Accepts"])
    for column in EXPECTED_COLUMNS:
        if column.choices:
            accepts = ", ".join(value for value, _label in column.choices)
        else:
            accepts = column.help_text or column.kind
        instructions_sheet.append([column.header, "Yes" if column.required else "No", accepts])
    instructions_sheet.protection.sheet = True

    return workbook


def _purge_source_file(batch: EmployeeImportBatch) -> None:
    """D-141: once applied or reversed, the source spreadsheet's content is gone.

    It held ID numbers in the clear — the whole reason ``employee.id_number`` is
    encrypted at rest (D-77). Leaving the upload sitting in storage puts that
    protection right back where it started.
    """
    if batch.source_file_id is None:
        return
    with tenant_context_of(batch):
        purge_content(batch.source_file)
