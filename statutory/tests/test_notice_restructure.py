"""BCCCI clause 21.1(b) in three rows, and the loader path that got it there (D-315).

r4 carried four bands with a boundary at six months on the probation band. The
gazette does say "between 4 weeks ... and six months" (page 36, verbatim in the
notes), but the six months never binds: clause 3 caps probation at four months,
and off probation the notice is two weeks on both sides of six months. So r5
states the same notice in three rows, loaded by ``--restructure``, which refuses
unless every reachable answer is unchanged.

PROVE EVERY GUARD FAILS: the restructure is watched refusing a changed answer,
a table other than notice bands, and a restructure with nothing to supersede.
"""

from __future__ import annotations

import datetime
import json
import pathlib
from decimal import Decimal

import pytest

from statutory import resolve
from statutory.loader import (
    ASKS_FOR_PROBATION,
    FIXTURE_ORDER,
    ReferenceDataLoadError,
    SupersedeRefusedError,
    load_reference_data,
    restructure_permits,
)
from statutory.models import (
    ReferenceDataVersion,
    Sector,
    SectorArea,
    TerminationNoticeBand,
)

pytestmark = pytest.mark.django_db

REFERENCE = pathlib.Path(__file__).resolve().parents[2] / "reference"
NOTICE = "ref-2026.04.01-bccci-notice.json"
R4 = pathlib.Path(__file__).resolve().parent / "fixtures" / "ref-2026.04.01-bccci-notice-r4.json"
APRIL = datetime.date(2026, 4, 1)


def read(path) -> dict:
    return json.loads(pathlib.Path(path).read_text(encoding="utf-8"))


def r5() -> dict:
    return read(REFERENCE / NOTICE)


@pytest.fixture
def with_r4(db):
    """Everything shipped, but the BCCCI notice bands as r4 had them — the
    state Kobus's database is in before the restructure."""
    for name in FIXTURE_ORDER:
        load_reference_data(read(R4) if name == NOTICE else read(REFERENCE / name))


@pytest.fixture
def shipped(db):
    for name in FIXTURE_ORDER:
        load_reference_data(read(REFERENCE / name))


def area_b():
    return (
        Sector.objects.get(code=Sector.Code.CONTRACT_CLEANING),
        SectorArea.objects.get(sector__code=Sector.Code.CONTRACT_CLEANING, code="AREA_B"),
    )


def notice(start, on, *, on_probation):
    sector, area = area_b()
    band = resolve.notice_band(
        sector, on, employment_start_date=start, sector_area=area, on_probation=on_probation
    )
    return band.notice_value.normalize(), band.notice_unit


def restructure(document):
    return load_reference_data(
        document,
        supersede="REF-2026.04.01-BCCCI-NOTICE-r4",
        reason="Clause 21.1(b)'s six-month boundary changes no answer (D-315)",
        restructure=True,
    )


# ================================================================ the load


def test_the_restructure_keeps_r4s_rows_and_puts_three_in_their_place(with_r4):
    report = restructure(r5())

    assert report.superseded == "REF-2026.04.01-BCCCI-NOTICE-r4"
    assert report.retired == {"termination_notice_band": 4}
    assert report.updated == {}, "nothing was re-encoded; four rows were retired"
    version = ReferenceDataVersion.objects.get(version_label="REF-2026.04.01-BCCCI-NOTICE-r5")
    assert version.supersedes.version_label == "REF-2026.04.01-BCCCI-NOTICE-r4"
    sector, area = area_b()
    # The 2026 agreement's rule set: the 2023 predecessor has its own four bands.
    bands = TerminationNoticeBand.objects.filter(
        termination_rule_set__sector=sector,
        termination_rule_set__sector_area=area,
        termination_rule_set__effective_from=APRIL,
    )
    kept = bands.filter(superseded_by_version=version)
    live = bands.filter(superseded_by_version__isnull=True)
    assert sorted(kept.values_list("sequence", flat=True)) == [1, 2, 3, 4]
    assert sorted(live.values_list("sequence", flat=True)) == [5, 6, 7]
    assert all("21.1(b)" in band.source_reference for band in live)


def test_the_verbatim_six_month_phrase_is_kept_in_the_citation_notes(shipped):
    sector, area = area_b()
    probation = TerminationNoticeBand.objects.get(
        termination_rule_set__sector=sector,
        termination_rule_set__sector_area=area,
        termination_rule_set__effective_from=APRIL,
        probation_condition="on_probation",
        superseded_by_version__isnull=True,
    )
    assert "between 4 weeks as in sub clause (i) above and six months" in probation.notes
    assert probation.service_to_value is None


# ======================================================== both sides of each edge


@pytest.mark.parametrize(
    "service_to, on_probation, expected",
    [
        # Exactly four weeks: the first band, "during the first four weeks".
        (datetime.timedelta(weeks=4), True, (Decimal("1"), "days")),
        (datetime.timedelta(weeks=4), False, (Decimal("1"), "days")),
        # Four weeks and a day: the lanes part.
        (datetime.timedelta(weeks=4, days=1), True, (Decimal("1"), "weeks")),
        (datetime.timedelta(weeks=4, days=1), False, (Decimal("2"), "weeks")),
    ],
)
def test_the_four_week_boundary(shipped, service_to, on_probation, expected):
    on = datetime.date(2026, 9, 30)
    assert notice(on - service_to, on, on_probation=on_probation) == expected


def test_exactly_four_months_on_the_last_day_of_probation_and_the_day_after(shipped):
    """Clause 3's maximum. On its last day the employee is still on probation
    (employees/probation.py); the day after, they are not."""
    start = datetime.date(2026, 5, 1)
    assert notice(start, datetime.date(2026, 9, 1), on_probation=True) == (Decimal("1"), "weeks")
    assert notice(start, datetime.date(2026, 9, 2), on_probation=False) == (Decimal("2"), "weeks")


@pytest.mark.parametrize("months", [6, 12])
def test_six_months_and_a_year_are_two_weeks(shipped, months):
    from dateutil.relativedelta import relativedelta

    on = datetime.date(2027, 6, 30)
    start = on - relativedelta(months=months)
    assert notice(start, on, on_probation=False) == (Decimal("2"), "weeks")
    assert notice(start, on + datetime.timedelta(days=1), on_probation=False) == (
        Decimal("2"),
        "weeks",
    )


def test_the_probation_cap_the_open_band_relies_on_is_at_most_six_months(shipped):
    """The open-ended probation band is right only while clause 3 keeps
    probation inside the clause's own six-month window. If the cap is ever
    loaded above six months this fails, and the band has to be split again."""
    sector, area = area_b()
    cap = resolve.parameter_or_none("PROBATION_MAX_MONTHS", APRIL, sector=sector, sector_area=area)
    assert cap is not None
    assert cap.value_numeric <= 6


# ====================================================== the guard, refusing


def test_a_restructure_that_changes_an_answer_is_refused_and_writes_nothing(with_r4):
    document = r5()
    document["tables"]["termination_notice_band"][2]["notice_value"] = "3"

    with pytest.raises(SupersedeRefusedError, match="answer\\(s\\) changed") as refused:
        restructure(document)

    assert "off probation: ('2', 'weeks') became ('3', 'weeks')" in str(refused.value)
    assert not ReferenceDataVersion.objects.filter(
        version_label="REF-2026.04.01-BCCCI-NOTICE-r5"
    ).exists()
    assert not TerminationNoticeBand.objects.filter(superseded_by_version__isnull=False).exists()


def test_a_restructure_carrying_another_table_is_refused(with_r4):
    document = r5()
    document["tables"]["public_holiday"] = []
    document["tables"]["public_holiday"] = [{"country_code": "ZA"}]

    with pytest.raises(SupersedeRefusedError, match="restructures termination_notice_band"):
        restructure(document)


def test_restructure_without_a_supersede_is_refused(with_r4):
    with pytest.raises(ReferenceDataLoadError, match="is a kind of --supersede"):
        load_reference_data(r5(), restructure=True)


def test_a_plain_supersede_still_refuses_the_same_file(with_r4):
    """The ordinary re-encoding path is unchanged: rows it has never seen are
    new data, and it says so."""
    with pytest.raises(SupersedeRefusedError, match="refused"):
        load_reference_data(r5(), supersede="REF-2026.04.01-BCCCI-NOTICE-r4", reason="test")


def test_without_the_probation_cap_the_same_restructure_is_refused(with_r4, monkeypatch):
    """PROVE EVERY GUARD FAILS, on the one judgement in the proof. The open
    band answers one week for "on probation at seven months" where r4 answered
    two — a state clause 3 makes unreachable. Take the cap away and that state
    is compared like any other, and the restructure is refused, as it must be."""
    monkeypatch.setattr("statutory.loader.PROBATION_CAP_PARAMETER", "NO_SUCH_PARAMETER")

    with pytest.raises(SupersedeRefusedError, match=r"on probation: \('2', 'weeks'\) became"):
        restructure(r5())


# ================================ the one answer the restructure may move (D-315)


def test_an_unstated_call_past_six_months_is_now_asked_for_probation(shipped):
    """r4 answered two weeks here without being told; r5's lanes both run past
    six months, so the resolver asks. No production caller is unstated -
    payroll/termination.py always says - and no figure moved."""
    sector, area = area_b()
    with pytest.raises(
        resolve.StatutoryValueMissingError,
        match="depends on whether they are still on probation",
    ):
        resolve.notice_band(
            sector,
            datetime.date(2027, 1, 5),
            employment_start_date=datetime.date(2026, 4, 1),
            sector_area=area,
        )


@pytest.mark.parametrize(
    "on_probation, was, now, permitted",
    [
        # The allowance: unstated, a figure becomes "asks".
        (None, ("2", "weeks"), ASKS_FOR_PROBATION, True),
        # Unstated may NOT become a different figure.
        (None, ("2", "weeks"), ("1", "weeks"), False),
        # Unstated may not go from refusing to answering, either.
        (None, ("refused", "x"), ("2", "weeks"), False),
        # A caller who SAYS may see nothing move at all.
        (False, ("2", "weeks"), ASKS_FOR_PROBATION, False),
        (True, ("1", "weeks"), ("2", "weeks"), False),
        # Unchanged is always fine.
        (True, ("1", "weeks"), ("1", "weeks"), True),
    ],
)
def test_what_the_restructure_permits(on_probation, was, now, permitted):
    assert restructure_permits(on_probation, was, now) is permitted


# ================================= the workbook, across a restructure (D-315)


def _retired_groups(path):
    """Mark every group on the 2026 agreement's r4 bands, as Kobus had."""
    from openpyxl import load_workbook

    sector, area = area_b()
    pks = {
        str(pk)
        for pk in TerminationNoticeBand.objects.filter(
            termination_rule_set__sector=sector,
            termination_rule_set__sector_area=area,
            termination_rule_set__effective_from=APRIL,
        ).values_list("pk", flat=True)
    }
    book = load_workbook(path)
    sheet = book["Checks"]
    marked = []
    for row in range(2, sheet.max_row + 1):
        key = sheet.cell(row=row, column=10).value or ""
        parts = key.split(":")
        if len(parts) == 4 and parts[1] == "termination_notice_band" and parts[2] in pks:
            sheet.cell(row=row, column=11, value="Y")
            sheet.cell(row=row, column=12, value="kobus@example.com")
            sheet.cell(row=row, column=13, value="2026-09-24")
            marked.append(key)
    book.save(path)
    return marked


@pytest.fixture
def kobus(db):
    from django.contrib.auth import get_user_model

    return get_user_model().objects.create_user(email="kobus@example.com", password="x" * 14)


def test_ticks_recorded_on_retired_bands_let_the_workbook_be_refreshed(with_r4, kobus, tmp_path):
    """The export's --force refuses over ticks nobody imported (D-272). Ticks on
    bands a restructure retired ARE imported - their rows are kept - so the
    refresh goes through. Found on Kobus's own workbook the first time."""
    from django.core.management import call_command
    from django.core.management.base import CommandError

    path = tmp_path / "workbook.xlsx"
    call_command("exportverification", str(path), verbosity=0)
    marked = _retired_groups(path)
    assert marked
    try:
        call_command("importverification", str(path), current_through="2027-02-28", verbosity=0)
    except CommandError as refused:
        # Verifying a VERSION needs every figure in it ticked, which this test
        # does not do; the ticks themselves are recorded regardless (D-256).
        assert "Every SOUND tick in the workbook was still recorded" in str(refused)
    from statutory.models import ReferenceFigureCheck

    assert ReferenceFigureCheck.objects.filter(
        row_key__startswith="termination_notice_band:"
    ).exists()
    restructure(r5())

    call_command("exportverification", str(path), force=True, verbosity=0)


def test_ticks_on_retired_bands_that_were_never_imported_still_refuse(with_r4, kobus, tmp_path):
    """PROVE EVERY GUARD FAILS: the allowance is for RECORDED ticks only."""
    from django.core.management import call_command
    from django.core.management.base import CommandError

    path = tmp_path / "workbook.xlsx"
    call_command("exportverification", str(path), verbosity=0)
    marked = _retired_groups(path)
    restructure(r5())

    with pytest.raises(
        CommandError, match=r"marked row\(s\) that are NOT in the database"
    ) as refused:
        call_command("exportverification", str(path), force=True, verbosity=0)
    assert marked[0] in str(refused.value)
