"""The import round trip, and the silence that hid it (D-272).

A workbook with seventeen marked rows imported and recorded NOTHING, and said
nothing about it. The pass this file protects is 156 checks by one person over
several evenings; an importer that can quietly discard an evening's work is the
one defect that makes the whole exercise not worth starting.

Everything here is end to end through the real commands, against a real
exported file. The round trip had never been run in one test before: the
pieces were each tested and the seam between them was not.
"""

from __future__ import annotations

import json
import pathlib

import pytest
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import CommandError
from openpyxl import load_workbook

from statutory.loader import load_reference_data
from statutory.models import ReferenceFigureCheck, Sector, SectorArea

pytestmark = [pytest.mark.django_db, pytest.mark.statutory]

REFERENCE = pathlib.Path(__file__).resolve().parents[2] / "reference"

BASE = ("ref-2026.03.01-sick-accrual.json", "ref-2026.03.01-employment.json")
#: Four figures in one version, so "some rows ticked" is a real state.
EXTRA = "ref-2026.04.01-bccci-leave-types.json"

CHECKED_COLUMN = 11
BY_COLUMN = 12
DATE_COLUMN = 13
KEY_COLUMN = 10


@pytest.fixture
def loader(db):
    return get_user_model().objects.create_user(email="loader@example.com", password="x" * 14)


@pytest.fixture
def kobus(db):
    return get_user_model().objects.create_user(email="kobus@example.com", password="x" * 14)


@pytest.fixture
def loaded(loader):
    for name in BASE:
        load_reference_data(
            json.loads((REFERENCE / name).read_text(encoding="utf-8")), loaded_by=loader
        )
    return loader


@pytest.fixture
def with_extra(loaded):
    sector = Sector.objects.create(
        code=Sector.Code.CONTRACT_CLEANING, name="Contract cleaning sector"
    )
    SectorArea.objects.create(
        sector=sector, code="AREA_B", name="Area B", uses_bargaining_council_rates=True
    )
    load_reference_data(
        json.loads((REFERENCE / EXTRA).read_text(encoding="utf-8")), loaded_by=loaded
    )
    return loaded


def export(path, **options):
    call_command("exportverification", str(path), verbosity=0, **options)
    return path


def mark(path, count, *, by="kobus@example.com", when="2026-09-20"):
    """Tick the first ``count`` rows the way a person with a gazette would."""
    book = load_workbook(path)
    sheet = book["Checks"]
    marked = []
    for row in range(2, sheet.max_row + 1):
        if len(marked) >= count:
            break
        key = sheet.cell(row=row, column=KEY_COLUMN).value
        if not key:
            continue
        sheet.cell(row=row, column=CHECKED_COLUMN, value="Y")
        sheet.cell(row=row, column=BY_COLUMN, value=by)
        sheet.cell(row=row, column=DATE_COLUMN, value=when)
        marked.append(key)
    book.save(path)
    return marked


def ticked(path):
    """Row keys the workbook shows as already checked."""
    sheet = load_workbook(path)["Checks"]
    return [
        sheet.cell(row=row, column=KEY_COLUMN).value
        for row in range(2, sheet.max_row + 1)
        if (sheet.cell(row=row, column=CHECKED_COLUMN).value or "").strip()
    ]


# ------------------------------------------------------------- THE ROUND TRIP


def test_two_marked_two_recorded_two_carried_forward(with_extra, kobus, tmp_path, capsys):
    """THE TEST THAT HAD NEVER BEEN RUN. Export, tick, import, re-export — the
    seam between the three commands, not each of them alone."""
    first = export(tmp_path / "one.xlsx")
    marked = mark(first, 2)
    assert len(marked) == 2

    call_command("importverification", str(first), current_through="2029-02-28", verbosity=1)

    recorded = ReferenceFigureCheck.objects.filter(
        outcome=ReferenceFigureCheck.Outcome.CHECKED
    ).count()
    assert recorded >= 2, "two ticked groups must record at least one check each"
    assert "Recorded" in capsys.readouterr().out

    second = export(tmp_path / "two.xlsx")
    assert sorted(ticked(second)) == sorted(marked), (
        "a fresh export must carry forward exactly what was imported, so an "
        "evening's work is never asked for twice"
    )


# --------------------------------------------- the cause of the silent zero


def test_a_tick_by_someone_with_no_account_is_refused_and_named(with_extra, tmp_path, capsys):
    """THE CAUSE (D-272). ``_sift`` checked that 'checked by' was NOT BLANK and
    nothing more; ``_record`` then looked the email up, found no user, and
    ``continue``d — with a comment claiming the row was "named in _sift's
    problems", which it never was. Seventeen rows, nothing written, nothing
    said.

    The email a person types is their own, and a verifier who has never been
    given a login is the ordinary case rather than a freak one.
    """
    path = export(tmp_path / "one.xlsx")
    mark(path, 2, by="nobody@example.com")

    with pytest.raises(CommandError):
        call_command("importverification", str(path), current_through="2029-02-28", verbosity=1)

    output = capsys.readouterr().out
    assert "nobody@example.com" in output, "name the address that has no account"
    assert ReferenceFigureCheck.objects.count() == 0


def test_a_tick_with_an_unreadable_date_is_refused_and_named(with_extra, kobus, tmp_path, capsys):
    """The second silent ``continue`` in the same loop, for the same reason."""
    path = export(tmp_path / "one.xlsx")
    mark(path, 2, when="last Tuesday")

    with pytest.raises(CommandError):
        call_command("importverification", str(path), current_through="2029-02-28", verbosity=1)

    output = capsys.readouterr().out
    assert "last Tuesday" in output
    assert ReferenceFigureCheck.objects.count() == 0


def test_the_command_can_never_be_silent_about_recording_nothing(with_extra, tmp_path, capsys):
    """Whatever else happens, a run that writes nothing SAYS nothing was
    written. The count used to print only ``if recorded``, so zero printed
    nothing at all and the run read as success."""
    path = export(tmp_path / "one.xlsx")

    call_command("importverification", str(path), current_through="2029-02-28", verbosity=1)

    output = capsys.readouterr().out
    assert "0 marked row(s) read, 0 recorded." in output, (
        "the count line is unconditional now; it used to print only when "
        "something was written, so nil printed nothing and read as success"
    )


def test_marked_rows_that_record_nothing_are_reported_as_such(with_extra, kobus, tmp_path, capsys):
    """Importing the same workbook twice records nothing the second time, which
    is correct — and must still say so, naming how many marked rows it read.
    "17 marked, 0 recorded" is the line whose absence cost an evening."""
    path = export(tmp_path / "one.xlsx")
    mark(path, 2)
    call_command("importverification", str(path), current_through="2029-02-28", verbosity=0)
    capsys.readouterr()

    call_command("importverification", str(path), current_through="2029-02-28", verbosity=1)

    output = capsys.readouterr().out
    assert "2 marked row(s) read, 0 recorded." in output
    assert "Nothing was written" in output, (
        "and it must say WHY nothing was written, or a correct no-op is "
        "indistinguishable from the silent discard this test exists for"
    )


# ------------------------------------------------------------ a stale workbook


def test_a_workbook_exported_before_the_data_changed_is_refused(loaded, kobus, tmp_path):
    """A workbook is a snapshot of the check groups. Load a new fixture and the
    groups it was built from are no longer the groups that exist, so what it
    does NOT mention is no longer "nothing to do" — it is "rows this file never
    knew about". Importing it silently would report a version complete when
    figures in it had never been looked at."""
    path = export(tmp_path / "one.xlsx")
    mark(path, 1)

    sector = Sector.objects.create(
        code=Sector.Code.CONTRACT_CLEANING, name="Contract cleaning sector"
    )
    SectorArea.objects.create(
        sector=sector, code="AREA_B", name="Area B", uses_bargaining_council_rates=True
    )
    load_reference_data(
        json.loads((REFERENCE / EXTRA).read_text(encoding="utf-8")), loaded_by=loaded
    )

    with pytest.raises(CommandError) as caught:
        call_command("importverification", str(path), current_through="2029-02-28", verbosity=0)

    assert "export a fresh" in str(caught.value).lower() or "stale" in str(caught.value).lower()


def test_a_workbook_that_matches_the_database_imports_normally(with_extra, kobus, tmp_path):
    """Watched NOT firing: the staleness guard must not refuse every workbook."""
    path = export(tmp_path / "one.xlsx")
    mark(path, 1)

    call_command("importverification", str(path), current_through="2029-02-28", verbosity=0)

    assert ReferenceFigureCheck.objects.exists()


# ----------------------------------------------------- --force and lost work


def test_force_refuses_to_overwrite_a_workbook_holding_unimported_ticks(
    with_extra, kobus, tmp_path
):
    """``--force`` exists for a file whose ticks are already in the database.
    Over a file holding work nobody has imported, it is the destructive
    operation this whole design was built to make impossible (D-256)."""
    path = export(tmp_path / "one.xlsx")
    mark(path, 2)

    with pytest.raises(CommandError) as caught:
        export(path, force=True)

    assert "not yet imported" in str(caught.value) or "importverification" in str(caught.value)


def test_force_overwrites_freely_once_the_ticks_are_imported(with_extra, kobus, tmp_path):
    """Watched NOT firing. Once the work is in the database the file is
    disposable, which is the whole point of recording ticks as evidence."""
    path = export(tmp_path / "one.xlsx")
    mark(path, 2)
    call_command("importverification", str(path), current_through="2029-02-28", verbosity=0)

    export(path, force=True)

    assert len(ticked(path)) == 2, "and the fresh file pre-fills what was imported"
