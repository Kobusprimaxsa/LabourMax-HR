"""The verification workbook, and the one thing it must never do.

P2's last 582 figures are a human job, and the workbook is the shape that job
has been missing. Two rules are load-bearing and both are tested here by
watching them work and then watching them refuse:

**Grouped by SOURCE DOCUMENT.** A person opens GN R.7083 once and checks every
figure citing it. Grouped by table, the same gazette is opened as many times as
there are tables citing it.

**The import records a verification and can never change a figure** (D-251). If
a value in the spreadsheet differs from the loaded one, the version is REFUSED
and the row is named. A convenient importer that wrote figures back is exactly
how a rate typed into Excel at 23:00 ends up on a payslip.
"""

from __future__ import annotations

import datetime
import io
import json
import pathlib

import pytest
from django.contrib.auth import get_user_model
from django.core.management import CommandError, call_command
from openpyxl import load_workbook

from statutory import verification
from statutory.loader import load_reference_data
from statutory.models import ReferenceDataVersion, ReferenceFigureCheck, StatutoryParameter

pytestmark = [pytest.mark.django_db, pytest.mark.statutory]

REFERENCE = pathlib.Path(__file__).resolve().parents[2] / "reference"

#: Two shipped fixtures with no foreign keys to anything, so a test database can
#: hold them alone. Real fixtures rather than invented ones, because
#: ``row_versions()`` recovers the version by replaying what is in reference/.
FIXTURES = ("ref-2026.03.01-sick-accrual.json", "ref-2026.03.01-employment.json")
SICK = "REF-2026.03.01-SICK-ACCRUAL"
EMPLOYMENT = "REF-2026.03.01-EMPLOYMENT"


@pytest.fixture
def loader(db):
    return get_user_model().objects.create_user(
        email="loader@example.com", password="x" * 14, first_name="A", last_name="Loader"
    )


@pytest.fixture
def checker(db):
    return get_user_model().objects.create_user(
        email="checker@example.com", password="x" * 14, first_name="A", last_name="Checker"
    )


@pytest.fixture
def loaded(loader):
    for name in FIXTURES:
        load_reference_data(
            json.loads((REFERENCE / name).read_text(encoding="utf-8")), loaded_by=loader
        )
    return loader


@pytest.fixture
def workbook(loaded, tmp_path):
    path = tmp_path / "verification.xlsx"
    call_command("exportverification", str(path), verbosity=0)
    return path


def figures(path):
    return load_workbook(path)["Checks"]


def mark_all(path, *, checked="Y", by="checker@example.com", when="2026-09-20"):
    """Tick the whole workbook the way a person would, and save it."""
    book = load_workbook(path)
    sheet = book["Checks"]
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row=row, column=11, value=checked)
        sheet.cell(row=row, column=12, value=by)
        sheet.cell(row=row, column=13, value=when)
    book.save(path)
    return path


# ------------------------------------------------- the grouping, which is the point


def test_the_workbook_is_grouped_by_source_document_and_not_by_table(loaded, tmp_path):
    """THE design decision (D-251). Every figure citing one document sits
    together, so the gazette is opened once. Grouped by table it would be opened
    once per table that cites it, and the job never gets finished."""
    lines = verification.lines()
    seen_order = [line.document for line in lines]

    assert seen_order == sorted(seen_order), "documents must not interleave"
    for document in set(seen_order):
        first, last = (
            seen_order.index(document),
            len(seen_order) - 1 - seen_order[::-1].index(document),
        )
        assert last - first + 1 == seen_order.count(document), (
            f"{document} is split across the workbook instead of being one block"
        )


@pytest.mark.parametrize(
    ("reference", "document", "clause"),
    [
        (
            "Basic Conditions of Employment Act 75 of 1997, s37(1)(a)",
            "Basic Conditions of Employment Act 75 of 1997",
            "s37(1)(a)",
        ),
        (
            # The FIRST marker wins: the last would invent a document called
            # "... s35, read with Form BCEA1A (Regulation 2)".
            "Basic Conditions of Employment Act 75 of 1997, s35, read with Form BCEA1A "
            "(Regulation 2), Summary of the Act",
            "Basic Conditions of Employment Act 75 of 1997",
            "s35, read with Form BCEA1A (Regulation 2), Summary of the Act",
        ),
        (
            # A marker inside brackets is not a marker, or the document name is
            # left with an unclosed parenthesis.
            "GN 5970, GG 52232, March 2025 (Basic Conditions of Employment Act 75 of 1997, s6(3))",
            "GN 5970, GG 52232, March 2025 (Basic Conditions of Employment Act 75 of 1997, s6(3))",
            "",
        ),
        (
            # SD1's own name contains the word "clauses".
            "Sectoral Determination 1: Contract Cleaning Sector, current consolidated text "
            "(clauses 3, 8-24), clause 23(1)(a)",
            "Sectoral Determination 1: Contract Cleaning Sector, current consolidated text "
            "(clauses 3, 8-24)",
            "clause 23(1)(a)",
        ),
        (
            # All five SD7 pinpoints have to land on one document.
            "Sectoral Determination 7: Domestic Worker Sector, published in Regulation Gazette "
            "No. 7434 (consolidated text), read with Basic Conditions of Employment Act 75 of "
            "1997 s37(1)(c); SD7 clause 24(1)(a)",
            "Sectoral Determination 7: Domestic Worker Sector, published in Regulation Gazette "
            "No. 7434 (consolidated text)",
            "read with Basic Conditions of Employment Act 75 of 1997 s37(1)(c); SD7 clause "
            "24(1)(a)",
        ),
        (
            "GN R.7083, GG 54075, 3 February 2026 (National Minimum Wage Act 9 of 2018)",
            "GN R.7083, GG 54075, 3 February 2026 (National Minimum Wage Act 9 of 2018)",
            "",
        ),
    ],
)
def test_a_citation_splits_into_its_document_and_its_pinpoint(reference, document, clause):
    assert verification.split_citation(reference) == (document, clause)


def test_a_figure_is_rendered_the_way_a_person_reads_it(loaded):
    """R 32,40 and not 32.4000. A verifier comparing against a gazette should
    not have to translate the database's precision in their head."""
    row = StatutoryParameter.objects.first()
    field = StatutoryParameter._meta.get_field("value_numeric")

    row.unit = StatutoryParameter.Unit.ZAR
    row.value_numeric = verification.decimal.Decimal("32.4000")
    assert verification.format_value(row, field) == "R 32,40"

    row.value_numeric = verification.decimal.Decimal("1234567.89")
    assert verification.format_value(row, field) == "R 1 234 567,89"

    row.unit = StatutoryParameter.Unit.PERCENT
    row.value_numeric = verification.decimal.Decimal("1.00")
    assert verification.format_value(row, field) == "1,00 %"


def test_the_workbook_says_it_is_not_a_source_of_truth(workbook):
    summary = load_workbook(workbook)["Summary"]
    text = " ".join(str(c.value) for row in summary.iter_rows() for c in row if c.value)
    assert "NEVER A SOURCE OF TRUTH" in text


def test_a_machine_verified_version_does_not_read_as_done(loaded, tmp_path, checker):
    """The four BCCCI versions were verified by a development identity. They must
    not look finished on the summary sheet or nobody will ever check them."""
    version = ReferenceDataVersion.objects.get(version_label=SICK)
    machine = get_user_model().objects.create_user(
        email="claude-verification@labourmax.invalid", password="x" * 14
    )
    version.verified_at = datetime.datetime(2026, 9, 19, tzinfo=datetime.UTC)
    version.verified_by_user = machine
    version.golden_tests_passed = True
    version.data_current_through = datetime.date(2029, 2, 28)
    version.save()

    path = tmp_path / "machine.xlsx"
    call_command("exportverification", str(path), verbosity=0)
    summary = load_workbook(path)["Summary"]
    text = " ".join(str(c.value) for row in summary.iter_rows() for c in row if c.value)

    assert "MACHINE ONLY" in text
    assert "still needs a human" in text


def test_exporting_over_an_existing_workbook_is_refused(workbook):
    """Overwriting one is throwing away somebody's evening — the ticks live only
    in the file."""
    with pytest.raises(CommandError) as raised:
        call_command("exportverification", str(workbook), verbosity=0)
    assert "already exists" in str(raised.value)


# ------------------------------------------------------------- 2f: the round trip


def test_a_fully_checked_workbook_verifies_its_versions_and_moves_no_figure(workbook, checker):
    """THE ROUND TRIP. Export, tick everything, import: the versions verify and
    every figure is exactly as it was."""
    before = {line.key: line.value for line in verification.lines()}
    assert before, "the fixture loaded nothing"

    mark_all(workbook)
    call_command(
        "importverification",
        str(workbook),
        current_through="2027-02-28",
        golden_tests_passed=True,
        verbosity=0,
    )

    for label in (SICK, EMPLOYMENT):
        version = ReferenceDataVersion.objects.get(version_label=label)
        assert version.verified_at is not None, label
        assert version.verified_by_user.email == "checker@example.com"
        assert version.data_current_through == datetime.date(2027, 2, 28)
        assert version.is_usable

    after = {line.key: line.value for line in verification.lines()}
    assert after == before, "importing a workbook must never move a figure"


def test_a_figure_edited_in_the_workbook_refuses_the_version_and_names_the_row(workbook, checker):
    """The failure this whole architecture exists to prevent, watched refusing.
    Somebody 'corrects' a rate in Excel; the import must not carry it into the
    reference data, and must not quietly verify around it either."""
    mark_all(workbook)
    book = load_workbook(workbook)
    sheet = book["Checks"]
    edited_key = sheet.cell(row=2, column=10).value
    original = sheet.cell(row=2, column=7).value
    sheet.cell(row=2, column=7, value="value R 99,99")
    book.save(workbook)

    with pytest.raises(CommandError) as raised:
        call_command(
            "importverification",
            str(workbook),
            current_through="2027-02-28",
            golden_tests_passed=True,
            verbosity=0,
        )

    assert "refused" in str(raised.value).lower()

    group = verification.group_by_key(edited_key)
    assert ReferenceDataVersion.objects.get(version_label=group.version).verified_at is None
    assert group.summary == original, "not one figure in the group may have moved"
    assert not ReferenceFigureCheck.objects.filter(row_key__in=group.row_keys).exists(), (
        "a group whose summary drifted was not checked against what is loaded, "
        "so none of its figures may be recorded"
    )


def test_the_refusal_names_the_row_and_both_values(workbook, checker, capsys):
    mark_all(workbook)
    book = load_workbook(workbook)
    sheet = book["Checks"]
    key = sheet.cell(row=2, column=10).value
    sheet.cell(row=2, column=7, value="value R 99,99")
    book.save(workbook)

    with pytest.raises(CommandError):
        call_command("importverification", str(workbook), current_through="2027-02-28")

    output = capsys.readouterr().out
    assert key in output
    assert "R 99,99" in output
    assert "NEVER writes a figure" in output


# --------------------------------------------------------------- 2d: QUERY refuses


def test_a_queried_row_refuses_the_version_and_repeats_the_note(workbook, checker):
    mark_all(workbook)
    book = load_workbook(workbook)
    sheet = book["Checks"]
    sheet.cell(row=2, column=11, value="QUERY")
    sheet.cell(row=2, column=14, value="Gazette says 27 days, not 26.")
    book.save(workbook)

    with pytest.raises(CommandError):
        call_command("importverification", str(workbook), current_through="2027-02-28")

    assert ReferenceDataVersion.objects.get(version_label=SICK).verified_at is None


def test_a_query_with_no_note_is_refused_for_saying_nothing(workbook, checker, capsys):
    mark_all(workbook)
    book = load_workbook(workbook)
    book["Checks"].cell(row=2, column=11, value="QUERY")
    book.save(workbook)

    with pytest.raises(CommandError):
        call_command("importverification", str(workbook), current_through="2027-02-28")

    assert "does not say what is wrong" in capsys.readouterr().out


# ------------------------------------------------- 2c and 2e: who, and running twice


def test_the_loader_may_not_verify_their_own_load(workbook, loaded):
    """The second-person rule, reached through the workbook rather than around
    it. verifystatutory refuses this and so must anything that calls it."""
    mark_all(workbook, by="loader@example.com")

    with pytest.raises(CommandError):
        call_command("importverification", str(workbook), current_through="2027-02-28")

    assert ReferenceDataVersion.objects.get(version_label=SICK).verified_at is None


def test_the_export_says_which_versions_a_named_verifier_may_not_sign(loaded, tmp_path):
    """2c: found out before the evening is spent, not after."""
    path = tmp_path / "for-loader.xlsx"
    call_command("exportverification", str(path), verifier="loader@example.com", verbosity=0)
    summary = load_workbook(path)["Summary"]
    text = " ".join(str(c.value) for row in summary.iter_rows() for c in row if c.value)

    assert "loader@example.com may verify?" in text
    assert "NO — you loaded it" in text


def test_a_half_finished_workbook_verifies_nothing_and_says_how_far_it_got(
    workbook, checker, capsys
):
    mark_all(workbook, checked="N")
    book = load_workbook(workbook)
    sheet = book["Checks"]
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=4).value == SICK:
            sheet.cell(row=row, column=11, value="Y")
    book.save(workbook)

    call_command("importverification", str(workbook), current_through="2027-02-28", verbosity=1)

    assert ReferenceDataVersion.objects.get(version_label=SICK).verified_at is not None
    assert ReferenceDataVersion.objects.get(version_label=EMPLOYMENT).verified_at is None
    assert "Still outstanding" in capsys.readouterr().out


def test_running_the_import_twice_leaves_the_first_run_alone(workbook, checker, capsys):
    """2e: four documents on Tuesday, five on Thursday. The second run must not
    complain about Tuesday's work or re-stamp it with Thursday's date."""
    mark_all(workbook)
    call_command(
        "importverification",
        str(workbook),
        current_through="2027-02-28",
        golden_tests_passed=True,
        verbosity=0,
    )
    first = ReferenceDataVersion.objects.get(version_label=SICK).verified_at

    call_command(
        "importverification",
        str(workbook),
        current_through="2028-02-29",
        golden_tests_passed=True,
        verbosity=1,
    )

    version = ReferenceDataVersion.objects.get(version_label=SICK)
    assert version.verified_at == first, "a recorded verification is not re-stamped"
    assert version.data_current_through == datetime.date(2027, 2, 28)
    assert "Already recorded" in capsys.readouterr().out


def test_a_dry_run_writes_nothing(workbook, checker):
    mark_all(workbook)
    call_command(
        "importverification",
        str(workbook),
        current_through="2027-02-28",
        golden_tests_passed=True,
        dry_run=True,
        verbosity=0,
    )
    assert ReferenceDataVersion.objects.get(version_label=SICK).verified_at is None


# ------------------------------ the ticks are evidence, so they live in the database


def mark_keys(path, keys, *, checked="Y", by="checker@example.com", when="2026-09-20", note=""):
    """Tick only the named row keys, leaving the rest blank."""
    book = load_workbook(path)
    sheet = book["Checks"]
    wanted = set(keys)
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=10).value in wanted:
            sheet.cell(row=row, column=11, value=checked)
            sheet.cell(row=row, column=12, value=by)
            sheet.cell(row=row, column=13, value=when)
            if note:
                sheet.cell(row=row, column=14, value=note)
    book.save(path)
    return path


def ticks(path) -> dict[str, str]:
    """What the workbook says is checked, keyed by row key."""
    sheet = load_workbook(path)["Checks"]
    return {
        sheet.cell(row=row, column=10).value: (sheet.cell(row=row, column=11).value or "")
        for row in range(2, sheet.max_row + 1)
    }


def test_a_re_export_after_loading_new_rows_carries_every_earlier_tick_forward(
    workbook, checker, tmp_path
):
    """THE HAZARD THIS CLOSES. Three P2 items are still to load, and each one
    needs a re-export. Before this, the first re-export threw away however many
    evenings had accumulated, because the ticks existed only in the file.

    So: tick half, import, load a reference row that did not exist when the
    workbook was written, re-export, and the earlier ticks must still be there
    with the new figures unchecked beside them.
    """
    keys = sorted(ticks(workbook))
    half = keys[: len(keys) // 2]
    assert half, "the fixture produced too few figures to halve"

    mark_keys(workbook, half)
    call_command("importverification", str(workbook), current_through="2027-02-28", verbosity=0)

    load_reference_data(
        json.loads((REFERENCE / "ref-2026.03.01-remuneration.json").read_text(encoding="utf-8")),
        loaded_by=get_user_model().objects.get(email="loader@example.com"),
    )

    reexported = tmp_path / "round-two.xlsx"
    call_command("exportverification", str(reexported), verbosity=0)
    after = ticks(reexported)

    assert set(after) > set(keys), "the new reference row did not reach the workbook"
    for key in half:
        assert after[key] == "Y", f"{key} lost its tick on re-export"
    for key in set(after) - set(keys):
        assert after[key] == "", f"{key} is new and must arrive unchecked"


@pytest.mark.parametrize("order", ["forward", "backward"])
def test_two_partial_workbooks_accumulate_the_same_way_in_either_order(
    loaded, checker, tmp_path, order
):
    """The one-person constraint becomes a convention rather than a data-loss
    risk: whoever ticks what, in whichever order the files are imported, the
    database ends up the same. Git still cannot merge the binary; nothing is
    lost when it tries.

    Both parameters assert the SAME expected state, which is what makes this a
    test of order-independence rather than two tests of one order each.
    """
    first, second = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    call_command("exportverification", str(first), verbosity=0)
    call_command("exportverification", str(second), verbosity=0)

    keys = sorted(ticks(first))
    assert len(keys) >= 2, "too few check groups to split between two people"
    mark_keys(first, keys[::2])
    mark_keys(second, keys[1::2])

    for path in (first, second) if order == "forward" else (second, first):
        call_command("importverification", str(path), current_through="2027-02-28", verbosity=0)

    state = {
        check.row_key: (check.outcome, check.checked_by.email)
        for check in ReferenceFigureCheck.objects.all()
    }

    expected = {
        figure_key: (ReferenceFigureCheck.Outcome.CHECKED, "checker@example.com")
        for key in keys
        for figure_key in verification.group_by_key(key).row_keys
    }
    assert state == expected, (
        "every FIGURE recorded exactly once - the grouping is presentation, the "
        "evidence is per figure - whoever ticked it and in whichever order"
    )
    assert len(expected) > len(keys), "the groups must actually cover several figures each"


def test_a_check_record_can_never_be_edited_or_deleted(workbook, checker):
    """Invariant 4 on this table, watched refusing. A tick is evidence: a change
    of mind inserts a second row and the first one stands."""
    from django.db import DatabaseError, transaction

    mark_all(workbook)
    call_command("importverification", str(workbook), current_through="2027-02-28", verbosity=0)
    check = ReferenceFigureCheck.objects.first()
    assert check is not None

    with pytest.raises(DatabaseError) as raised, transaction.atomic():
        check.note = "actually, no"
        check.save(update_fields=["note"])
    assert "append-only" in str(raised.value)

    with pytest.raises(DatabaseError) as raised, transaction.atomic():
        ReferenceFigureCheck.objects.filter(pk=check.pk).delete()
    assert "append-only" in str(raised.value)


def test_changing_your_mind_inserts_a_second_row_and_keeps_the_first(workbook, checker):
    """A figure queried on Tuesday and checked on Thursday reads in order,
    rather than being overwritten into a single reassuring Y."""
    key = sorted(ticks(workbook))[0]
    mark_keys(workbook, [key], checked="QUERY", note="Gazette may say 27, not 26.")
    with pytest.raises(CommandError):
        # The QUERY blocks its version, and is recorded all the same - a query
        # nobody wrote down is a question asked twice.
        call_command("importverification", str(workbook), current_through="2027-02-28", verbosity=0)

    mark_keys(workbook, [key], checked="Y", when="2026-09-22", note="Checked again; 26 is right.")
    call_command("importverification", str(workbook), current_through="2027-02-28", verbosity=0)

    figure = verification.group_by_key(key).row_keys[0]
    history = list(
        ReferenceFigureCheck.objects.filter(row_key=figure).order_by("recorded_at", "id")
    )
    assert [c.outcome for c in history] == [
        ReferenceFigureCheck.Outcome.QUERIED,
        ReferenceFigureCheck.Outcome.CHECKED,
    ]
    assert verification.latest_checks()[figure].outcome == ReferenceFigureCheck.Outcome.CHECKED


def test_the_summary_counts_are_right_before_anyone_opens_the_file(workbook, checker, tmp_path):
    """Counted from the DATABASE. A COUNTIFS would read as zero on a freshly
    written file, because openpyxl writes the formula and only Excel evaluates
    it — and zero is exactly the wrong answer to "how far have I got"."""
    keys = sorted(ticks(workbook))
    mark_keys(workbook, keys[:1])
    call_command("importverification", str(workbook), current_through="2027-02-28", verbosity=0)

    fresh = tmp_path / "fresh.xlsx"
    call_command("exportverification", str(fresh), verbosity=0)

    summary = load_workbook(fresh, data_only=True)["Summary"]
    totals = [row for row in summary.iter_rows(values_only=True) if row and row[0] == "TOTAL"]
    assert totals, "the summary has no total row"
    checks_total, figures_total, done, queried, outstanding = totals[0][1:6]
    assert checks_total == len(keys)
    assert figures_total == len(verification.lines())
    assert done == 1, "the one imported tick must show without Excel recalculating"
    assert queried == 0
    assert outstanding == len(keys) - 1


# ------------------------------------------- D-258: one lookup, one tick


def test_a_group_covers_one_row_of_one_table_and_never_spans_two(loaded):
    """The clause column alone is not granular enough anywhere it matters — the
    SARS guide cites no pinpoint at all, the Public Holidays Act cites Schedule 1
    for all of them, and a rule set cites a seven-clause range for 24 figures. So
    the database row is part of the key, or a tick would cover figures the person
    never looked at."""
    for group in verification.check_groups():
        rows = {key.rsplit(":", 1)[0] for key in group.row_keys}
        assert len(rows) == 1, f"{group.key} spans {rows}"
        assert len({(f.document, f.clause) for f in group.figures}) == 1


def test_no_group_is_bigger_than_the_cap(loaded):
    """Above the cap the inline summary wraps to three lines in Excel and the
    tick stops meaning "I read all of these", which is what makes a group safe."""
    for group in verification.check_groups():
        assert 1 <= len(group.figures) <= verification.GROUP_CAP, group.key


def test_every_figure_belongs_to_exactly_one_group(loaded):
    """Nothing may fall between two groups and go unchecked, and nothing may be
    counted twice."""
    lines = verification.lines()
    seen = [key for group in verification.check_groups(lines) for key in group.row_keys]

    assert sorted(seen) == sorted(line.key for line in lines)
    assert len(seen) == len(set(seen))


def test_a_split_row_says_which_part_it_is(db):
    """A rule set is 22 to 26 figures and splits into parts. The person has to
    know the clause reference is the whole row's, not the part's - the split is
    by cap and NOT by clause, because the citation does not say which figure
    came from which clause."""
    from statutory.loader import FIXTURE_ORDER

    for name in FIXTURE_ORDER:
        path = REFERENCE / name
        if path.exists():
            load_reference_data(json.loads(path.read_text(encoding="utf-8")))

    split = [g for g in verification.check_groups() if g.parts > 1]
    assert split, "the rule set rows are above the cap and must split"
    for group in split:
        assert f"part {group.part} of {group.parts}" in group.label


def test_the_group_row_shows_every_figure_it_covers(loaded):
    """A tick means "all of these match". It can only mean that if all of them
    are on the row."""
    for group in verification.check_groups():
        for figure in group.figures:
            assert figure.value in group.summary
            assert figure.figure_label in group.summary


def test_ticking_one_group_records_one_check_per_figure(workbook, checker):
    """D-258's load-bearing sentence: the grouping is presentation, the evidence
    stays per figure."""
    key = sorted(ticks(workbook))[0]
    group = verification.group_by_key(key)
    assert len(group.figures) > 1, "pick a fixture whose first group covers several figures"

    mark_keys(workbook, [key])
    call_command("importverification", str(workbook), current_through="2027-02-28", verbosity=0)

    recorded = set(ReferenceFigureCheck.objects.values_list("row_key", flat=True))
    assert recorded == set(group.row_keys)


# ---------------- two people may verify one version between them (D-270)


@pytest.fixture
def second_checker(db):
    return get_user_model().objects.create_user(
        email="second@example.com", password="x" * 14, first_name="B", last_name="Checker"
    )


#: Four figures in one version, so it can actually be split between two people.
#: Every UNSCOPED fixture in reference/ holds exactly one parameter row, so this
#: one needs its sector and area created first - which is the point of using a
#: real fixture rather than an invented one (``row_versions()`` recovers the
#: version by replaying what is in reference/).
SPLITTABLE_FIXTURE = "ref-2026.04.01-bccci-leave-types.json"
SPLITTABLE = "REF-2026.04.01-BCCCI-LEAVE-TYPES"


@pytest.fixture
def split_workbook(loaded, tmp_path):
    from statutory.models import Sector, SectorArea

    sector = Sector.objects.create(
        code=Sector.Code.CONTRACT_CLEANING, name="Contract cleaning sector"
    )
    SectorArea.objects.create(
        sector=sector, code="AREA_B", name="Area B", uses_bargaining_council_rates=True
    )
    load_reference_data(
        json.loads((REFERENCE / SPLITTABLE_FIXTURE).read_text(encoding="utf-8")),
        loaded_by=loaded,
    )
    path = tmp_path / "split.xlsx"
    call_command("exportverification", str(path), verbosity=0)
    return path


def keys_for(version):
    """Every check-group key belonging to one version, in workbook order."""
    return [group.key for group in verification.check_groups() if group.version == version]


def test_two_people_splitting_one_version_verify_it_between_them(
    split_workbook, checker, second_checker
):
    """THE CASE THE WORKBOOK'S OWN SHAPE PRODUCES. It is organised by SOURCE
    DOCUMENT because that is how a person verifies — one gazette, one sitting —
    and versions cut ACROSS documents. So splitting the pass by document
    between two people guarantees that some version is checked by both, and
    refusing that made the obvious way of sharing the work impossible.

    Two people checking different figures is a STRONGER result than one person
    checking all of them, not a weaker one.
    """
    half = keys_for(SPLITTABLE)
    assert len(half) >= 4, "this test needs a version with something to split"
    mark_keys(split_workbook, half[:1], by="checker@example.com")
    mark_keys(split_workbook, half[1:], by="second@example.com")

    call_command(
        "importverification", str(split_workbook), current_through="2027-02-28", verbosity=0
    )

    version = ReferenceDataVersion.objects.get(version_label=SPLITTABLE)
    assert version.verified_at is not None, "a version checked by two people is still verified"
    assert {user.email for user in version.verified_by_users.all()} == {
        "checker@example.com",
        "second@example.com",
    }, "every checker is recorded, not just whoever signed"
    assert version.verified_by_user.email in {"checker@example.com", "second@example.com"}


def test_the_loader_may_not_verify_their_own_load_even_as_the_second_checker(
    split_workbook, loaded, second_checker, capsys
):
    """THE RULE THE ONE-CHECKER REFUSAL WAS ACCIDENTALLY COVERING. The loader
    check ran against ``checkers[0]`` only, so with the refusal lifted a loader
    could verify their own load simply by being the second name on it. It runs
    over EVERY checker now."""
    half = keys_for(SPLITTABLE)
    mark_keys(split_workbook, half[:1], by="second@example.com")
    mark_keys(split_workbook, half[1:], by="loader@example.com")

    with pytest.raises(CommandError):
        call_command(
            "importverification", str(split_workbook), current_through="2027-02-28", verbosity=1
        )

    output = capsys.readouterr().out
    assert "loaded this version and may not verify it" in output, (
        "it must refuse for the RIGHT reason - before this change the same call "
        "was refused merely for having two checkers, which would have made this "
        "test pass while the loader rule leaked"
    )
    assert "loader@example.com" in output
    assert ReferenceDataVersion.objects.get(version_label=SPLITTABLE).verified_at is None


def test_a_machine_among_the_checkers_taints_the_whole_version(split_workbook, checker):
    """D-262 held only because there was one verifier to look at. With several,
    a version is machine-verified if ANY of them is a development identity —
    otherwise half a version checked by nobody rides in on the other half."""
    machine = get_user_model().objects.create_user(
        email="claude-verification@labourmax.invalid", password="x" * 14
    )
    half = keys_for(SPLITTABLE)
    mark_keys(split_workbook, half[:1], by="checker@example.com")
    mark_keys(split_workbook, half[1:], by=machine.email)

    call_command(
        "importverification", str(split_workbook), current_through="2027-02-28", verbosity=0
    )

    version = ReferenceDataVersion.objects.get(version_label=SPLITTABLE)
    assert version.is_machine_verified is True
    assert version.is_usable is False
    assert ReferenceDataVersion.in_force_on(datetime.date(2026, 6, 1)) is None, (
        "unusable_q() reaches the OTHER checkers through an EXISTS subquery, and "
        "this is the exclude() path that would silently stop excluding if it did not"
    )


def test_a_dry_run_reports_what_it_would_record_not_what_the_database_holds(
    workbook, checker, capsys
):
    """It read the DATABASE, so however much was ticked it said "0 of N
    checked" — which reads as a failure and sent a person hunting a bug that
    was not there. Preview IS apply, rolled back (D-145), so the report is
    accurate by construction rather than by a second code path that has to be
    kept in step."""
    mark_all(workbook)

    call_command(
        "importverification",
        str(workbook),
        current_through="2027-02-28",
        dry_run=True,
        verbosity=1,
    )

    output = capsys.readouterr().out
    assert "DRY RUN" in output
    assert SICK in output
    assert "0 of" not in output, "a fully ticked workbook must not report nothing checked"
    assert "Would verify" in output or "Verified" in output

    assert ReferenceDataVersion.objects.get(version_label=SICK).verified_at is None
    assert ReferenceFigureCheck.objects.count() == 0, "a dry run still writes nothing"


# ------------------------------------- the Checked column's vocabulary (D-282)
#
# One list, in statutory/verification.py, read by the dropdown, by the --force
# guard and by the importer. It used to be two dicts in two command modules
# plus a third value the dropdown offered that neither mentioned — N, meaning
# "not yet". The importer dropped it in silence, correctly, and the export then
# refused to overwrite the file because it held a mark with no record behind
# it, naming "delete it by hand" as the way out. The one value meaning "I have
# not done this" was the only one that made the workbook un-refreshable.


def mark_one(path, row, value, *, by="checker@example.com", when="2026-09-20"):
    book = load_workbook(path)
    sheet = book["Checks"]
    sheet.cell(row=row, column=11, value=value)
    sheet.cell(row=row, column=12, value=by)
    sheet.cell(row=row, column=13, value=when)
    book.save(path)
    return path


def test_the_dropdown_offers_exactly_what_the_importer_records(workbook):
    """No N. The list is built from ``RECORDED_AS``, so a value can only be
    offered if something reads it back."""
    validations = load_workbook(workbook)["Checks"].data_validations.dataValidation
    formulas = [v.formula1 for v in validations]

    assert '"Y,QUERY"' in formulas
    assert not any("N" in f.strip('"').split(",") for f in formulas), formulas


def test_the_two_halves_of_the_vocabulary_are_one_list(loaded):
    """``CHECKED_TEXT`` is derived from ``RECORDED_AS`` rather than written out
    beside it, so they cannot drift — which is how N came to exist."""
    assert verification.CHECKED_TEXT == {
        outcome: mark for mark, outcome in verification.RECORDED_AS.items()
    }
    for mark in verification.RECORDED_AS:
        assert verification.CHECKED_TEXT[verification.RECORDED_AS[mark]] == mark


# ---------------------------------------------- watching the --force guard REFUSE


def test_force_refuses_over_a_tick_nobody_imported(workbook):
    """THE GUARD, unchanged and still the point (D-272). A Y in the file and
    nothing in the database is an evening of checking that --force would
    destroy."""
    mark_one(workbook, 2, "Y")

    with pytest.raises(CommandError) as raised:
        call_command("exportverification", str(workbook), force=True, verbosity=0)

    assert "would discard them" in str(raised.value)


def test_force_refuses_over_an_unimported_query_too(workbook):
    """A QUERY is recorded evidence — it is the whole value of one."""
    mark_one(workbook, 2, "QUERY")

    with pytest.raises(CommandError) as raised:
        call_command("exportverification", str(workbook), force=True, verbosity=0)

    assert "would discard them" in str(raised.value)


# --------------------------------------------- watching it NOT fire over an N


@pytest.mark.parametrize("mark", ["N", "n", " n "])
def test_force_is_not_blocked_by_a_deferral(workbook, mark):
    """Nothing was ever recorded for an N, so there is nothing to discard. The
    lower-case cases are the second half of the same bug: the guard compared
    raw text while the importer upper-cased, so the two disagreed about what a
    cell said."""
    mark_one(workbook, 2, mark)

    call_command("exportverification", str(workbook), force=True, verbosity=0)

    assert not str(figures(workbook).cell(row=2, column=11).value or "").strip(), (
        "the row should come back blank, which is what 'not yet' means"
    )


def test_a_deferral_that_was_passed_over_is_reported(workbook):
    """Said out loud, not quietly dropped — D-272's own lesson. A row that
    stops being marked without a word is how somebody loses track of an
    evening."""
    mark_one(workbook, 2, "N")
    out = io.StringIO()

    call_command("exportverification", str(workbook), force=True, stdout=out, verbosity=1)

    assert "1 row(s) in the overwritten file carried a mark that records nothing" in out.getvalue()


def test_force_over_an_imported_workbook_says_nothing_about_deferrals(workbook, checker):
    """Watched NOT firing. The ordinary case — everything ticked and imported —
    must not grow a line about rows nobody deferred."""
    mark_all(workbook)
    call_command("importverification", str(workbook), current_through="2027-02-28", verbosity=0)
    out = io.StringIO()

    call_command("exportverification", str(workbook), force=True, stdout=out, verbosity=1)

    assert "records nothing" not in out.getvalue()


def test_an_n_from_an_older_workbook_still_imports_as_not_yet(workbook, checker):
    """A file exported before the dropdown changed still has N cells in it.
    They must go on meaning 'not yet' rather than becoming an error."""
    mark_all(workbook)
    mark_one(workbook, 2, "N")

    call_command("importverification", str(workbook), current_through="2027-02-28", verbosity=0)

    key = figures(workbook).cell(row=2, column=10).value
    group = verification.group_by_key(key)
    recorded = set(verification.latest_checks())
    assert not any(figure.key in recorded for figure in group.figures), (
        "an N must record nothing at all"
    )
